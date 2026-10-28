"""07: Cek silang Excel (Q14). NG rate dan ranking supplier untuk dua bulan dihitung dengan formula Excel
(INDEX/MATCH untuk supplier tiap pemakaian, SUMIFS untuk jumlah, COUNTIF untuk ranking) dari data bersih,
lalu dihitung ulang oleh LibreOffice (headless) dan dibandingkan dengan view SQL.
Hasil: outputs/crosscheck.xlsx (berisi nilai terhitung) dan results/2-analysis/crosscheck_excel.csv.
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import openpyxl
import pandas as pd
from openpyxl.styles import Font
import common as C

CHECK_MONTHS = ["2026-06", "2026-08"]
lot = pd.read_csv(C.CLEAN / "material_lot_clean.csv")
use = pd.read_csv(C.CLEAN / "lot_usage_clean.csv", parse_dates=["date"])
suppliers = sorted(lot.supplier_code.unique())
sql = pd.read_csv(C.RES2 / "ng_month.csv")

wb = openpyxl.Workbook()
ws_lot = wb.active
ws_lot.title = "lot"
ws_lot.append(["lot_code", "supplier_code"])
for r in lot.sort_values("lot_code").itertuples():
    ws_lot.append([r.lot_code, r.supplier_code])
n_lot = len(lot) + 1

ws_u = wb.create_sheet("pemakaian")
ws_u.append(["tanggal", "lot_code", "input_qty", "ng_qty", "supplier", "bulan"])
u = use.sort_values(["date", "lot_code", "station_code", "model_code", "input_qty", "ng_material_qty"]).reset_index(drop=True)
for i, r in u.iterrows():
    row = i + 2
    ws_u.append([r.date.to_pydatetime(), r.lot_code, int(r.input_qty), int(r.ng_material_qty),
                 f"=INDEX(lot!$B$2:$B${n_lot},MATCH(B{row},lot!$A$2:$A${n_lot},0))",
                 f"=YEAR(A{row})*100+MONTH(A{row})"])
    ws_u.cell(row=row, column=1).number_format = "yyyy-mm-dd"
n_u = len(u) + 1

ws = wb.create_sheet("cek", 0)
ws.append(["bulan (yyyymm)", "supplier", "unit", "NG", "NG rate", "ranking (1 = NG rate tertinggi)"])
for c in ws[1]:
    c.font = Font(bold=True)
row = 2
for m in CHECK_MONTHS:
    first = row
    key = int(m.replace("-", ""))
    for s in suppliers:
        ws.append([key, s,
                   f"=SUMIFS(pemakaian!$C$2:$C${n_u},pemakaian!$E$2:$E${n_u},B{row},pemakaian!$F$2:$F${n_u},A{row})",
                   f"=SUMIFS(pemakaian!$D$2:$D${n_u},pemakaian!$E$2:$E${n_u},B{row},pemakaian!$F$2:$F${n_u},A{row})",
                   f"=D{row}/C{row}", None])
        row += 1
    last = row - 1
    for r in range(first, last + 1):
        ws.cell(row=r, column=6, value=f'=1+COUNTIF($E${first}:$E${last},">"&E{r})')
ws.column_dimensions["F"].width = 32

tmp = Path(tempfile.mkdtemp())
src = tmp / "crosscheck_raw.xlsx"
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)],
               check=True, capture_output=True, timeout=300)
calc = tmp / "out" / src.name
shutil.copy(calc, C.OUTPUTS / "crosscheck.xlsx")

vals = openpyxl.load_workbook(calc, data_only=True)["cek"]
rows = []
for r in vals.iter_rows(min_row=2, values_only=True):
    key, s, units, ng, rate, rank = r[:6]
    m = f"{str(key)[:4]}-{str(key)[4:]}"
    q = sql[(sql.month == m) & (sql.supplier_code == s)].iloc[0]
    rows.append(dict(month=m, supplier_code=s, units_excel=int(units), units_sql=int(q.units), ng_excel=int(ng), ng_sql=int(q.ng),
                     rate_excel=round(float(rate), 8), rate_sql=round(float(q.ng_rate), 8), rank_excel=int(rank), rank_sql=int(q.rank_ng)))
res = pd.DataFrame(rows)
res["cocok"] = ((res.units_excel == res.units_sql) & (res.ng_excel == res.ng_sql) &
                ((res.rate_excel - res.rate_sql).abs() < 1e-7) & (res.rank_excel == res.rank_sql)).astype(int)
res.to_csv(C.RES2 / "crosscheck_excel.csv", index=False)
print(res.to_string(index=False))
print(f"cocok: {res.cocok.sum()}/{len(res)}")
assert res.cocok.sum() == len(res), "Excel dan SQL tidak cocok"
